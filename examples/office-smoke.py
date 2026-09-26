"""Offline round-trip example for DSH's bundled Python and Office libraries.

Run with the Python path returned by load_workspace_dependencies:
    python office-smoke.py --output /path/to/a/new/folder --node /path/to/node
The output folder must be empty. No downloads, accounts or API keys are used.
"""
import argparse
import bz2
import ctypes
import hashlib
import importlib.metadata
import json
import lzma
from pathlib import Path
import sqlite3
import ssl
import subprocess
import sys
import zlib


def run(output, node=None):
    from PIL import Image, ImageDraw
    from docx import Document
    from docx.shared import Inches as DocxInches
    from lxml import etree
    from openpyxl import load_workbook
    from pptx import Presentation
    from pptx.util import Inches
    from pypdf import PdfReader
    from reportlab.pdfgen import canvas
    import xlsxwriter

    output = output.resolve()
    output.mkdir(parents=True, exist_ok=True)
    if any(output.iterdir()):
        raise ValueError("Use an empty output folder; existing files are never replaced.")
    checks = []
    payload = "DSH 中文路径 round trip".encode("utf-8")
    for module in (zlib, bz2, lzma):
        assert module.decompress(module.compress(payload)) == payload
    assert ctypes.sizeof(ctypes.c_void_p) == 8
    assert ssl.create_default_context().verify_mode == ssl.CERT_REQUIRED
    with sqlite3.connect(":memory:") as db:
        assert db.execute("select 6 * 7").fetchone()[0] == 42
    checks.append("stdlib: ssl, hashlib, sqlite3, ctypes, zlib, bz2, lzma")

    rows = [("小明", 88, 92, 90), ("小红", 95, 89, 94), ("小华", 76, 85, 82)]
    headers = ["姓名", "语文", "数学", "英语", "总分"]
    totals = [sum(row[1:]) for row in rows]
    data = output / "成绩数据.json"
    data.write_text(json.dumps({"rows": rows, "totals": totals}, ensure_ascii=False, indent=2), encoding="utf-8")

    picture = output / "示例图.png"
    image = Image.new("RGB", (640, 360), "white")
    draw = ImageDraw.Draw(image)
    draw.text((25, 20), "DSH Office library test", fill="black")
    for i, total in enumerate(totals):
        y = 75 + i * 80
        draw.rectangle((80, y, 80 + total, y + 40), fill=(45, 115, 205))
        draw.text((15, y + 12), f"S{i + 1}", fill="black")
        draw.text((95 + total, y + 12), str(total), fill="black")
    image.save(picture)
    with Image.open(picture) as loaded:
        loaded.load()
        assert loaded.size == (640, 360) and loaded.mode == "RGB"
    checks.append("Pillow: PNG write/read")

    spreadsheet = output / "成绩表.xlsx"
    with xlsxwriter.Workbook(str(spreadsheet)) as workbook:
        sheet = workbook.add_worksheet("成绩")
        sheet.write_row(0, 0, headers)
        sheet.set_column(0, 4, 14)
        for index, row in enumerate(rows, start=1):
            sheet.write_row(index, 0, row)
            sheet.write_formula(index, 4, f"=SUM(B{index + 1}:D{index + 1})", None, totals[index - 1])
    workbook = load_workbook(spreadsheet, data_only=True)
    assert [workbook["成绩"].cell(i, 5).value for i in range(2, 5)] == totals
    workbook.close()
    workbook = load_workbook(spreadsheet)
    assert workbook["成绩"]["E2"].value == "=SUM(B2:D2)"
    workbook.create_sheet("说明")["A1"] = "DSH 内置 openpyxl 修改成功；总分公式保留。"
    edited = output / "成绩表-已编辑.xlsx"
    workbook.save(edited)
    workbook.close()
    workbook = load_workbook(edited)
    assert workbook["说明"]["A1"].value.startswith("DSH")
    workbook.close()
    checks.append("XlsxWriter/openpyxl: XLSX formulas, cached totals and editable copy")

    document_path = output / "成绩报告.docx"
    document = Document()
    document.add_heading("学生成绩报告", 0)
    document.add_paragraph("文字、表格和图片均为独立对象，可以继续编辑。")
    table = document.add_table(rows=1, cols=5)
    for cell, heading in zip(table.rows[0].cells, headers):
        cell.text = heading
    for row, total in zip(rows, totals):
        for cell, value in zip(table.add_row().cells, [*row, total]):
            cell.text = str(value)
    document.add_picture(str(picture), width=DocxInches(5))
    document.save(document_path)
    loaded_document = Document(document_path)
    assert loaded_document.paragraphs[0].text == "学生成绩报告"
    assert loaded_document.tables[0].cell(1, 4).text == str(totals[0])
    assert len(loaded_document.inline_shapes) == 1
    checks.append("python-docx: DOCX text/table/image write/read")

    slides_path = output / "成绩演示.pptx"
    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[6])
    slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(9), Inches(0.7)).text = "学生成绩汇报（可编辑文字）"
    table = slide.shapes.add_table(4, 5, Inches(0.5), Inches(1.5), Inches(9), Inches(2)).table
    for j, heading in enumerate(headers):
        table.cell(0, j).text = heading
    for i, (row, total) in enumerate(zip(rows, totals), start=1):
        for j, value in enumerate([*row, total]):
            table.cell(i, j).text = str(value)
    slide = slides.slides.add_slide(slides.slide_layouts[6])
    slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(9), Inches(0.7)).text = "成绩图示（独立图片对象）"
    slide.shapes.add_picture(str(picture), Inches(1), Inches(1.5), width=Inches(8))
    slides.save(slides_path)
    loaded_slides = Presentation(slides_path)
    assert len(loaded_slides.slides) == 2
    assert loaded_slides.slides[0].shapes[0].has_text_frame
    assert loaded_slides.slides[0].shapes[1].has_table
    assert loaded_slides.slides[0].shapes[1].table.cell(1, 4).text == str(totals[0])
    assert loaded_slides.slides[1].shapes[1].image.size == (640, 360)
    checks.append("python-pptx: two slides with editable text/table and separate image")

    pdf_path = output / "成绩摘要.pdf"
    pdf = canvas.Canvas(str(pdf_path))
    pdf.setTitle("DSH Student Scores")
    pdf.drawString(60, 790, "DSH Office library test")
    for i, total in enumerate(totals):
        pdf.drawString(60, 750 - i * 25, f"Student S{i + 1}: total {total}")
    pdf.drawImage(str(picture), 60, 330, width=480, height=270)
    pdf.save()
    loaded_pdf = PdfReader(pdf_path)
    assert len(loaded_pdf.pages) == 1
    assert "Student S1: total 270" in loaded_pdf.pages[0].extract_text()
    checks.append("ReportLab/pypdf: PDF create, read and extract text")
    assert etree.fromstring(etree.tostring(etree.Element("office", verified="true"))).get("verified") == "true"
    checks.append("lxml: native XML parser/serializer")

    node_result = None
    if node:
        code = "const fs=require('node:fs');const d=JSON.parse(fs.readFileSync(process.argv[1],'utf8'));console.log(JSON.stringify({node:process.version,totals:d.rows.map(r=>r.slice(1).reduce((a,b)=>a+b,0))}));"
        result = subprocess.run([str(Path(node).resolve()), "-e", code, str(data)], check=True, capture_output=True, text=True, encoding="utf-8")
        node_result = json.loads(result.stdout)
        assert node_result["totals"] == totals
        checks.append("Node.js: independent process reads Unicode JSON path and checks totals")

    names = ["openpyxl", "python-docx", "python-pptx", "Pillow", "lxml", "XlsxWriter", "pypdf", "reportlab", "et_xmlfile", "typing_extensions", "charset-normalizer"]
    report = {"passed": True, "python": sys.executable, "pythonVersion": sys.version.split()[0],
              "node": node_result, "packages": {name: importlib.metadata.version(name) for name in names},
              "checks": checks, "files": [{"name": p.name, "bytes": p.stat().st_size, "sha256": hashlib.sha256(p.read_bytes()).hexdigest()} for p in sorted(output.iterdir())]}
    (output / "测试报告.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=True))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", required=True, type=Path, help="New or empty output directory")
    parser.add_argument("--node", help="Optional bundled Node executable")
    args = parser.parse_args()
    run(args.output, args.node)
