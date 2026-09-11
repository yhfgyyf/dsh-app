#!/usr/bin/env python3
"""Regenerate small document-preview fixtures; GUI tests use the committed files.

Requires python-docx, python-pptx and openpyxl only when regenerating fixtures.
"""

import argparse
import hashlib
import io
import json
import math
from pathlib import Path
import re
import struct
import wave
import xml.etree.ElementTree as ET
import zipfile
import zlib

from docx import Document
from docx.shared import Inches as WordInches
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.util import Inches, Pt


TEXT = "桌面文本预览：你好，世界。\nTXT_PREVIEW_OK\n第二行：中文编码验证。\n"
SHEET_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
FIXED_TIME = b"2000-01-01T00:00:00Z"


def picture_bytes():
    """A tiny RGB bar chart; no external media or imaging dependency."""
    width, height = 64, 48
    rows = []
    bars = [(8, 20, 22, (34, 96, 186)), (27, 39, 10, (12, 150, 140)), (46, 58, 28, (230, 130, 40))]
    for y in range(height):
        row = bytearray([0])
        for x in range(width):
            color = (255, 255, 255)
            for left, right, top, fill in bars:
                if left <= x < right and top <= y < 42:
                    color = fill
            row.extend(color)
        rows.append(row)

    def chunk(kind, data):
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF)

    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)) + chunk(b"IDAT", zlib.compress(b"".join(rows), 9)) + chunk(b"IEND", b"")


def stable_office(data, formula_caches=None):
    """Normalize ZIP timestamps/core dates and populate Excel formula caches."""
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(output, "w") as target:
        for name in sorted(source.namelist()):
            content = source.read(name)
            if name == "docProps/core.xml":
                content = re.sub(
                    rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*(</dcterms:(?:created|modified)>)",
                    lambda match: match[1] + FIXED_TIME + match[2],
                    content,
                )
            if formula_caches and name in formula_caches:
                ET.register_namespace("", SHEET_NS)
                root = ET.fromstring(content)
                for coordinate, cached in formula_caches[name].items():
                    cell = root.find(f".//{{{SHEET_NS}}}c[@r='{coordinate}']")
                    assert cell is not None and cell.find(f"{{{SHEET_NS}}}f") is not None
                    value = cell.find(f"{{{SHEET_NS}}}v")
                    if value is None:
                        value = ET.SubElement(cell, f"{{{SHEET_NS}}}v")
                    value.text = str(cached)
                content = ET.tostring(root, encoding="utf-8")
            entry = zipfile.ZipInfo(name, (2000, 1, 1, 0, 0, 0))
            entry.compress_type = zipfile.ZIP_DEFLATED
            entry.external_attr = 0o644 << 16
            target.writestr(entry, content, compresslevel=9)
    return output.getvalue()


def word_fixture(picture):
    document = Document()
    document.core_properties.title = "DSH 文档预览测试"
    document.core_properties.author = "DSH test fixture"
    document.add_heading("桌面文档预览测试", level=1)
    document.add_paragraph("DOCX 中文段落：你好，世界。预览应保留图片和表格。")
    document.add_paragraph("DOCX_PREVIEW_OK")
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ["项目", "结果"]):
        cell.text = text
    for values in [["中文表格", "通过"], ["图片数量", "1"]]:
        for cell, text in zip(table.add_row().cells, values):
            cell.text = text
    document.add_picture(io.BytesIO(picture), width=WordInches(1.5))
    data = io.BytesIO()
    document.save(data)
    return stable_office(data.getvalue())


def presentation_fixture(picture):
    presentation = Presentation()
    presentation.slide_width = Inches(10)
    presentation.slide_height = Inches(5.625)
    presentation.core_properties.title = "DSH 幻灯片预览测试"
    presentation.core_properties.author = "DSH test fixture"
    for number, title, sentence, shape_type in [
        (1, "演示文稿第一页", "中文文本：图片与形状预览", MSO_SHAPE.ROUNDED_RECTANGLE),
        (2, "演示文稿第二页", "第二页中文内容", MSO_SHAPE.OVAL),
    ]:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.35), Inches(9), Inches(0.7))
        title_box.text_frame.text = title
        title_box.text_frame.paragraphs[0].font.size = Pt(30)
        body = slide.shapes.add_textbox(Inches(0.5), Inches(1.3), Inches(6.2), Inches(1.4))
        body.text_frame.text = sentence + "\n" + f"PPTX_SLIDE_{number}_OK"
        for paragraph in body.text_frame.paragraphs:
            paragraph.font.size = Pt(22)
        shape = slide.shapes.add_shape(shape_type, Inches(0.7), Inches(3), Inches(2.7), Inches(1.6))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(34, 96, 186) if number == 1 else RGBColor(230, 130, 40)
        shape.text_frame.text = "矩形" if number == 1 else "椭圆"
        slide.shapes.add_picture(io.BytesIO(picture), Inches(6.8), Inches(2.8), width=Inches(2.5))
    data = io.BytesIO()
    presentation.save(data)
    return stable_office(data.getvalue())


def spreadsheet_fixture():
    workbook = Workbook()
    workbook.properties.title = "DSH 工作簿预览测试"
    workbook.properties.creator = "DSH test fixture"
    sales = workbook.active
    sales.title = "销售数据"
    for row in [
        ["产品", "单价", "数量", "金额"],
        ["测试商品", 12.5, 4, "=B2*C2"],
        ["第二商品", 8, 3, "=B3*C3"],
        ["合计", None, None, "=SUM(D2:D3)"],
    ]:
        sales.append(row)
    for cell in sales[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2260BA")
    sales.column_dimensions["A"].width = 18
    summary = workbook.create_sheet("汇总")
    summary.append(["项目", "数值"])
    summary.append(["中文汇总", "='销售数据'!D4"])
    summary.append(["校验标记", "XLSX_SECOND_SHEET_OK"])
    data = io.BytesIO()
    workbook.save(data)
    return stable_office(data.getvalue(), {
        "xl/worksheets/sheet1.xml": {"D2": 50, "D3": 24, "D4": 74},
        "xl/worksheets/sheet2.xml": {"B2": 74},
    })


def audio_fixture():
    output = io.BytesIO()
    with wave.open(output, "wb") as audio:
        audio.setnchannels(1)
        audio.setsampwidth(2)
        audio.setframerate(8000)
        frames = (struct.pack("<h", round(3000 * math.sin(2 * math.pi * 440 * index / 8000))) for index in range(2000))
        audio.writeframes(b"".join(frames))
    return output.getvalue()


def verify(directory):
    document = Document(directory / "sample.docx")
    assert "DOCX_PREVIEW_OK" in [paragraph.text for paragraph in document.paragraphs]
    assert document.tables[0].cell(1, 0).text == "中文表格"
    assert len(document.inline_shapes) == 1
    presentation = Presentation(directory / "sample.pptx")
    assert len(presentation.slides) == 2
    for index, slide in enumerate(presentation.slides, 1):
        text = "\n".join(shape.text_frame.text for shape in slide.shapes if shape.has_text_frame)
        assert f"PPTX_SLIDE_{index}_OK" in text
        assert sum(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in slide.shapes) == 1
    formulas = load_workbook(directory / "sample.xlsx", data_only=False)
    values = load_workbook(directory / "sample.xlsx", data_only=True)
    assert values.sheetnames == ["销售数据", "汇总"]
    assert formulas["销售数据"]["D4"].value == "=SUM(D2:D3)"
    assert [values["销售数据"][cell].value for cell in ["D2", "D3", "D4"]] == [50, 24, 74]
    assert values["汇总"]["B2"].value == 74
    assert values["汇总"]["B3"].value == "XLSX_SECOND_SHEET_OK"
    for name, encoding in [("utf8.txt", "utf-8"), ("utf8-bom.txt", "utf-8-sig"), ("utf16le.txt", "utf-16"), ("utf16be.txt", "utf-16"), ("gbk.txt", "gbk")]:
        assert (directory / name).read_bytes().decode(encoding) == TEXT
    with wave.open(str(directory / "sample.wav")) as audio:
        assert audio.getnframes() == 2000 and audio.getframerate() == 8000


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=Path(__file__).resolve().parent)
    arguments = parser.parse_args()
    directory = arguments.output
    directory.mkdir(parents=True, exist_ok=True)
    picture = picture_bytes()
    files = {
        "sample.png": picture,
        "sample.docx": word_fixture(picture),
        "sample.pptx": presentation_fixture(picture),
        "sample.xlsx": spreadsheet_fixture(),
        "sample.wav": audio_fixture(),
        "utf8.txt": TEXT.encode("utf-8"),
        "utf8-bom.txt": b"\xef\xbb\xbf" + TEXT.encode("utf-8"),
        "utf16le.txt": b"\xff\xfe" + TEXT.encode("utf-16-le"),
        "utf16be.txt": b"\xfe\xff" + TEXT.encode("utf-16-be"),
        "gbk.txt": TEXT.encode("gbk"),
    }
    for name, content in files.items():
        (directory / name).write_bytes(content)
    verify(directory)
    manifest = {
        "files": {name: {"bytes": len(content), "sha256": hashlib.sha256(content).hexdigest()} for name, content in files.items()},
        "expected": {
            "sample.docx": {"text": ["桌面文档预览测试", "DOCX_PREVIEW_OK", "中文表格", "通过"], "tables": 1, "images": 1},
            "sample.pptx": {"slides": 2, "text": ["演示文稿第一页", "PPTX_SLIDE_1_OK", "演示文稿第二页", "PPTX_SLIDE_2_OK"], "imagesPerSlide": 1},
            "sample.xlsx": {"sheets": ["销售数据", "汇总"], "cachedValues": {"销售数据!D2": 50, "销售数据!D3": 24, "销售数据!D4": 74, "汇总!B2": 74}, "text": ["测试商品", "XLSX_SECOND_SHEET_OK"]},
            "sample.wav": {"durationSeconds": 0.25, "sampleRate": 8000, "channels": 1},
            "textFiles": {"text": TEXT, "encodings": {"utf8.txt": "UTF-8", "utf8-bom.txt": "UTF-8 BOM", "utf16le.txt": "UTF-16LE BOM", "utf16be.txt": "UTF-16BE BOM", "gbk.txt": "GBK"}},
        },
    }
    (directory / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Verified {len(files)} fixtures ({sum(len(content) for content in files.values())} bytes) in {directory}")


if __name__ == "__main__":
    main()
